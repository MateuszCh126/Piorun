"""Eksport notatek wykladowych do DOCX (opcjonalny, obok HTML).

Mapuje notes_structured.json na dokument Word: naglowki sekcji, osadzone
slajdy i markdown notatek (naglowki, pogrubienia, listy). Uzywa python-docx,
ktory juz jest zaleznoscia projektu (tools/docx_builder.py).
"""
import json
import os

try:
    from core.config import get_settings
except ModuleNotFoundError:
    import sys

    sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
    from core.config import get_settings

SETTINGS = get_settings()


def _add_runs_with_bold(paragraph, text):
    """Dodaje tekst do akapitu, zamieniajac **...** na pogrubione fragmenty."""
    for i, part in enumerate(text.split("**")):
        if not part:
            continue
        run = paragraph.add_run(part)
        if i % 2 == 1:
            run.bold = True


def _render_markdown_to_docx(doc, md):
    """Minimalny renderer markdownu notatek: ## / ### / '- ' / **bold**."""
    for raw in str(md or "").split("\n"):
        line = raw.strip()
        if not line:
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=4)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=3)
        elif line.startswith("- "):
            _add_runs_with_bold(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            _add_runs_with_bold(doc.add_paragraph(), line)


def build_lecture_docx(session_dir, subject):
    """Buduje DOCX z notes_structured.json. Zwraca sciezke lub None (brak notatek)."""
    from docx import Document
    from docx.shared import Inches

    notes_json = os.path.join(session_dir, "notes_structured.json")
    if not os.path.exists(notes_json):
        return None
    with open(notes_json, "r", encoding="utf-8") as f:
        data = json.load(f)

    date_str = os.path.basename(os.path.normpath(session_dir)).split("_")[0]
    doc = Document()
    doc.add_heading(f"{subject} — wyklad {date_str}", 0)

    for chunk in data:
        if not isinstance(chunk, dict):
            continue
        start = int(chunk.get("time_start", 0))
        doc.add_heading(f"[{start // 60:02d}:{start % 60:02d}]", level=2)

        slide = chunk.get("slide")
        if slide:
            slide_path = os.path.join(session_dir, slide)
            if os.path.exists(slide_path):
                try:
                    doc.add_picture(slide_path, width=Inches(5.5))
                except Exception:
                    pass

        _render_markdown_to_docx(doc, chunk.get("notes_markdown", ""))

    obj_dir = os.path.join(str(SETTINGS.notes_root), subject)
    os.makedirs(obj_dir, exist_ok=True)
    out_path = os.path.join(obj_dir, f"{date_str}.docx")
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 3:
        print("Uzycie: python docx_export.py [sciezka_do_sesji] [temat]")
    else:
        print(build_lecture_docx(sys.argv[1], sys.argv[2]))
