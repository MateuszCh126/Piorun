"""Test integracyjny pipeline'u notatek: transcript -> notatki -> HTML -> DOCX.

LLM jest zamockowany (canned markdown), wiec test nie potrzebuje Whispera,
modelu Gemma ani sieci. Sprawdza, ze cztery moduly lacza sie poprawnie.
"""
import json

import pytest
from PIL import Image

import core.html_builder as hb
import tools.docx_export as dx
import tools.note_generator as ng


CANNED_MD = (
    "## Drzewa B+\n"
    "**Indeks zgrupowany**: dane fizycznie ułożone wg klucza.\n"
    "- pierwszy punkt\n"
    "- drugi punkt\n"
    "### Podsumowanie tej części\n"
    "Krótkie podsumowanie."
)


@pytest.fixture
def fake_session(tmp_path):
    session_dir = tmp_path / "2026-04-14_2031_BazyDanych"
    slides_dir = session_dir / "slides"
    slides_dir.mkdir(parents=True)

    transcript = [
        {"start": 0, "end": 100, "text": "Wprowadzenie do drzew."},
        {"start": 100, "end": 250, "text": "Definicja B plus."},
        {"start": 250, "end": 400, "text": "Indeksy w bazach."},
        {"start": 400, "end": 500, "text": "Podsumowanie."},
    ]
    (session_dir / "transcript.json").write_text(json.dumps(transcript), encoding="utf-8")

    timeline = [
        {"timestamp": 0, "file": "slides/s0.jpg"},
        {"timestamp": 260, "file": "slides/s1.jpg"},
    ]
    (session_dir / "slides_timeline.json").write_text(json.dumps(timeline), encoding="utf-8")
    for name in ("s0.jpg", "s1.jpg"):
        Image.new("RGB", (200, 120), (30, 30, 30)).save(slides_dir / name, "JPEG")

    return str(session_dir)


@pytest.fixture(autouse=True)
def mock_llm(monkeypatch):
    monkeypatch.setattr(ng.brain, "get_system_prompt", lambda: "SYS")
    monkeypatch.setattr(
        ng.brain,
        "execute_brain_loop",
        lambda history, *a, **k: list(history) + [{"role": "assistant", "content": CANNED_MD}],
    )


def test_full_pipeline(fake_session):
    # 1. Transkrypcja -> notatki (LLM zamockowany)
    notes_path = ng.generate_academic_notes(fake_session, "BazyDanych")
    with open(notes_path, encoding="utf-8") as f:
        notes = json.load(f)
    assert len(notes) == 2  # dwa chunki (granica > 300 s)
    assert "Drzewa B+" in notes[0]["notes_markdown"]
    assert notes[0]["slide"] == "slides/s0.jpg"
    assert notes[1]["slide"] == "slides/s1.jpg"

    # 2. HTML
    html_path = hb.build_lecture_page(fake_session, "BazyDanych")
    assert html_path
    html = open(html_path, encoding="utf-8").read()
    assert "BazyDanych" in html
    assert "<strong>Indeks zgrupowany</strong>" in html
    assert "data:image" in html  # slajd osadzony w base64

    # 3. DOCX
    docx_path = dx.build_lecture_docx(fake_session, "BazyDanych")
    assert docx_path and docx_path.endswith(".docx")
    import os
    assert os.path.getsize(docx_path) > 5000  # niepusty dokument z obrazami


def test_docx_bold_and_headings(tmp_path):
    from docx import Document

    session_dir = tmp_path / "2026-05-01_1000_Algo"
    session_dir.mkdir()
    notes = [{"chunk_index": 0, "time_start": 0, "time_end": 300, "slide": None,
              "notes_markdown": CANNED_MD}]
    (session_dir / "notes_structured.json").write_text(json.dumps(notes), encoding="utf-8")

    out = dx.build_lecture_docx(str(session_dir), "Algo")
    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    assert any("Indeks zgrupowany" in t for t in texts)
    # pogrubiony fragment istnieje
    bolds = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert any("Indeks zgrupowany" in b for b in bolds)
